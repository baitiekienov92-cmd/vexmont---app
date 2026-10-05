import io
import os
import re
import secrets
from datetime import datetime, timedelta, timezone
from functools import wraps

from flask import (Flask, abort, flash, g, redirect, render_template, request,
                   send_file, session, url_for)
from werkzeug.security import check_password_hash, generate_password_hash

import db
from stages import MEASURE_KINDS, N, TPL, stage_state

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "dev-change-me")
app.config["MAX_CONTENT_LENGTH"] = 40 * 1024 * 1024
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=60)
app.teardown_appcontext(db.close)
db.init()

ROLES = {"admin": "Генеральный директор", "devdir": "Директор по развитию", "pm": "Руководитель проекта",
         "supply": "Снабжение", "nach": "Начальник участка", "prorab": "Прораб", "okk": "Инженер ОКК",
         "lawyer": "Юрист", "accountant": "Бухгалтер", "client": "Заказчик"}
PR_ROLES = {"admin", "pm", "nach", "prorab"}          # сдают работы, вызывают ОКК
OKK_ROLES = {"admin", "okk"}                           # принимают работы
LEGAL_ROLES = {"admin", "devdir", "lawyer", "accountant"}
MEASURE_ROLES = {"admin", "pm", "supply", "nach"}      # назначают и закрывают замеры
CATS = {"rough": "Черновые материалы", "finish": "Чистовые материалы", "works": "Работы", "other": "Прочее"}
CONTRACT_KINDS = {"client": "С заказчиками", "contractor": "С подрядчиками", "supplier": "С поставщиками"}
ROLE_SHORT = {"admin": "ген. директор", "devdir": "дир. по развитию", "pm": "руководитель проекта", "supply": "снабжение",
              "nach": "начальник участка", "prorab": "прораб", "okk": "ОКК", "lawyer": "юрист",
              "accountant": "бухгалтер", "client": "заказчик"}
KZ = timezone(timedelta(hours=5))


def now():
    return datetime.now(KZ).strftime("%Y-%m-%d %H:%M:%S")


def ru_date(s):
    return f"{s[8:10]}.{s[5:7]}.{s[0:4]}" if s and len(s) >= 10 else ""


def pct(v):
    return f"{round(v, 1):.1f}".replace(".", ",") + "%"


app.jinja_env.filters["d"] = ru_date
app.jinja_env.filters["pct"] = pct
def money(v):
    return f"{(v or 0):,.0f}".replace(",", " ") + " ₸"


app.jinja_env.filters["money"] = money
app.jinja_env.globals.update(TPL=TPL, N=N, ROLES=ROLES, CATS=CATS, CONTRACT_KINDS=CONTRACT_KINDS,
                             MEASURE_KINDS=MEASURE_KINDS, PR_ROLES=PR_ROLES, OKK_ROLES=OKK_ROLES,
                             LEGAL_ROLES=LEGAL_ROLES, MEASURE_ROLES=MEASURE_ROLES)


def norm_phone(p):
    d = re.sub(r"\D", "", p or "")
    if len(d) == 11 and d[0] == "8":
        d = "7" + d[1:]
    if len(d) == 10:
        d = "7" + d
    return "+" + d if len(d) == 11 else ""


# ---------- CSRF ----------
def csrf_token():
    if "csrf" not in session:
        session["csrf"] = secrets.token_hex(16)
    return session["csrf"]


app.jinja_env.globals["csrf"] = csrf_token


@app.before_request
def load_user():
    g.user = None
    uid = session.get("uid")
    if uid:
        g.user = db.q("SELECT * FROM users WHERE id=?", (uid,), one=True)
        if not g.user or g.user["status"] == "blocked":
            session.pop("uid", None)
            g.user = None
    if request.method == "POST" and request.form.get("csrf") != session.get("csrf"):
        abort(400, "Сессия устарела, обновите страницу")


def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if not g.user:
            return redirect(url_for("login"))
        if g.user["status"] != "active":
            return redirect(url_for("pending"))
        return f(*a, **k)
    return w


def admin_required(f):
    @wraps(f)
    @login_required
    def w(*a, **k):
        if g.user["role"] != "admin":
            abort(403)
        return f(*a, **k)
    return w


def can(action):
    r = g.user["role"]
    return {"pr": r in PR_ROLES, "okk": r in OKK_ROLES}[action]


def is_client():
    return g.user["role"] == "client"


def is_gd():
    return g.user["role"] == "admin"


def staff_required(f):
    @wraps(f)
    @login_required
    def w(*a, **k):
        if is_client():
            abort(403)
        return f(*a, **k)
    return w


def sched_access():
    return is_gd() or bool(g.user.get("can_sched"))


app.jinja_env.globals.update(sched_access=lambda: g.user and sched_access())


# ---------- доступ к объектам ----------
def complex_ids():
    if g.user["role"] == "admin":
        return [r["id"] for r in db.q("SELECT id FROM complexes")]
    if is_client():
        return [r["complex_id"] for r in db.q("SELECT DISTINCT f.complex_id FROM flat_access a JOIN flats f ON f.id=a.flat_id WHERE a.user_id=?", (g.user["id"],))]
    return [r["complex_id"] for r in db.q("SELECT complex_id FROM access WHERE user_id=?", (g.user["id"],))]


def client_flat_ids():
    return {r["flat_id"] for r in db.q("SELECT flat_id FROM flat_access WHERE user_id=?", (g.user["id"],))}


def check_complex(cid):
    if cid not in complex_ids():
        abort(404)


def flat_or_404(fid):
    f = db.q("SELECT * FROM flats WHERE id=?", (fid,), one=True)
    if not f:
        abort(404)
    check_complex(f["complex_id"])
    if is_client() and f["id"] not in client_flat_ids():
        abort(404)
    return f


def in_list(ids):
    return "(" + ",".join(str(int(i)) for i in ids) + ")" if ids else "(0)"


# ---------- проценты ----------
def flat_stats(work_rows):
    """work_rows: строки works одной квартиры → (ОКК %, прораб %, принято, сдано, состояния этапов)."""
    by = {}
    for w in work_rows:
        by.setdefault(w["stage"], []).append(w["status"])
    states = [stage_state(by.get(i, [])) for i in range(N)]
    o = sum(1 for s in states if s == "ok")
    p = sum(1 for i in range(N) if by.get(i) and all(x in ("ok", "rv") for x in by[i]))
    sent = [bool(by.get(i)) and all(x in ("ok", "rv") for x in by[i]) for i in range(N)]
    return {"okk": o * 100 / N, "pr": p * 100 / N, "o": o, "p": p, "states": states, "sent": sent}


def complex_stats(cid):
    flats = db.q("SELECT * FROM flats WHERE complex_id=? ORDER BY id", (cid,))
    if g.get("user") and is_client():
        mine = client_flat_ids()
        flats = [f for f in flats if f["id"] in mine]
    works = db.q("SELECT w.flat_id, w.stage, w.status FROM works w JOIN flats f ON f.id=w.flat_id WHERE f.complex_id=?", (cid,))
    by = {}
    for w in works:
        by.setdefault(w["flat_id"], []).append(w)
    out = []
    for f in flats:
        out.append({**f, **flat_stats(by.get(f["id"], []))})
    n = len(out) or 1
    return out, {"okk": sum(f["okk"] for f in out) / n, "pr": sum(f["pr"] for f in out) / n}


# ---------- шапка: задачи и уведомления ----------
def task_works():
    ids = complex_ids()
    base = ("SELECT w.*, f.number, f.complex_id, c.name AS cname FROM works w JOIN flats f ON f.id=w.flat_id "
            f"JOIN complexes c ON c.id=f.complex_id WHERE f.complex_id IN {in_list(ids)} ")
    r = g.user["role"]
    if r == "okk":
        return db.q(base + "AND w.status='rv' ORDER BY w.date")
    if r == "admin":
        return db.q(base + "AND (w.status='rv' OR (w.status='wk' AND w.ret<>'')) ORDER BY w.date")
    if r in PR_ROLES:
        return db.q(base + "AND w.status='wk' AND w.ret<>'' ORDER BY w.date")
    return []


@app.context_processor
def header_counts():
    if not g.get("user") or g.user["status"] != "active":
        return {}
    ids = complex_ids()
    extra = f" AND f.id IN {in_list(client_flat_ids())}" if is_client() else ""
    unread = db.q(
        "SELECT COUNT(*) AS n FROM events e JOIN flats f ON f.id=e.flat_id "
        f"WHERE f.complex_id IN {in_list(ids)}{extra} AND e.created>? AND COALESCE(e.user_id,0)<>?",
        (g.user["notes_seen"] or "", g.user["id"]), one=True)["n"]
    newreq = db.q("SELECT COUNT(*) AS n FROM requests WHERE status='new'", one=True)["n"] if is_gd() else 0
    return {"tasks_n": len(task_works()), "unread": unread, "newreq": newreq}


def log(flat_id, stage, kind, text):
    db.x("INSERT INTO events(flat_id,stage,kind,text,user_id,created) VALUES(?,?,?,?,?,?)",
         (flat_id, stage, kind, text, g.user["id"], now()))


# ---------- регистрация и вход ----------
@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        name = request.form.get("name", "").strip()[:80]
        phone = norm_phone(request.form.get("phone"))
        pw = request.form.get("password", "")
        role = request.form.get("role", "prorab")
        company = request.form.get("company", "").strip()[:120]
        if not name or not phone:
            flash("Укажите имя и номер телефона в формате +7 7XX XXX XX XX")
        elif len(pw) < 6:
            flash("Пароль — минимум 6 символов")
        elif role not in ROLES or role == "admin":
            flash("Выберите роль")
        elif db.q("SELECT id FROM users WHERE phone=?", (phone,), one=True):
            flash("Этот номер уже зарегистрирован — войдите")
        else:
            first = db.q("SELECT COUNT(*) AS n FROM users", one=True)["n"] == 0
            uid = db.x("INSERT INTO users(name,phone,pw,role,status,company,created) VALUES(?,?,?,?,?,?,?)",
                       (name, phone, generate_password_hash(pw), "admin" if first else role,
                        "active" if first else "pending", company, now()), returning=True)
            db.commit()
            session.clear()
            session["uid"] = uid
            session.permanent = True
            if first:
                flash("Вы первый пользователь — назначены администратором")
                return redirect(url_for("admin_objects"))
            return redirect(url_for("pending"))
    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        phone = norm_phone(request.form.get("phone"))
        u = db.q("SELECT * FROM users WHERE phone=?", (phone,), one=True)
        if not u or not check_password_hash(u["pw"], request.form.get("password", "")):
            flash("Неверный номер или пароль")
        elif u["status"] == "blocked":
            flash("Доступ заблокирован администратором")
        else:
            session.clear()
            session["uid"] = u["id"]
            session.permanent = True
            return redirect(url_for("home"))
    return render_template("login.html")


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/pending")
def pending():
    if not g.user:
        return redirect(url_for("login"))
    if g.user["status"] == "active":
        return redirect(url_for("home"))
    return render_template("pending.html")


# ---------- квартиры ----------
@app.route("/")
@login_required
def home():
    ids = complex_ids()
    cs = db.q(f"SELECT * FROM complexes WHERE id IN {in_list(ids)} ORDER BY id")
    cards = []
    for c in cs:
        flats, tot = complex_stats(c["id"])
        cards.append({**c, "n": len(flats), **tot})
    return render_template("home.html", cards=cards, nav="home")


@app.route("/c/<int:cid>")
@login_required
def flats(cid):
    check_complex(cid)
    c = db.q("SELECT * FROM complexes WHERE id=?", (cid,), one=True)
    fl, tot = complex_stats(cid)
    return render_template("flats.html", c=c, flats=fl, tot=tot, nav="home")


@app.route("/f/<int:fid>")
@login_required
def flat(fid):
    f = flat_or_404(fid)
    c = db.q("SELECT * FROM complexes WHERE id=?", (f["complex_id"],), one=True)
    works = db.q("SELECT * FROM works WHERE flat_id=? ORDER BY stage, idx", (fid,))
    st = flat_stats(works)
    photos = db.q("SELECT p.id, p.work_id FROM photos p JOIN works w ON w.id=p.work_id WHERE w.flat_id=? ORDER BY p.id", (fid,))
    pics = {}
    for p in photos:
        pics.setdefault(p["work_id"], []).append(p["id"])
    by_stage = {}
    for w in works:
        by_stage.setdefault(w["stage"], []).append(w)
    hist = {}
    for e in db.q("SELECT e.*, u.name AS uname FROM events e LEFT JOIN users u ON u.id=e.user_id WHERE e.flat_id=? ORDER BY e.id DESC", (fid,)):
        hist.setdefault(e["stage"], []).append(e)
    users = {u["id"]: u for u in db.q("SELECT id, name, role FROM users")}
    cards = []
    for i, (name, tw, photo_only) in enumerate(TPL):
        ws = by_stage.get(i, [])
        state = st["states"][i]
        callers = [users[w["sent_by"]] for w in ws if w["sent_by"] in users and w["status"] in ("rv", "ok")]
        returned = any(w["status"] == "wk" and w["ret"] for w in ws)
        cards.append({"i": i, "name": name, "works": ws, "photo_only": photo_only, "state": state,
                      "acc": sum(1 for w in ws if w["status"] == "ok"), "total": len(ws), "returned": returned,
                      "caller": ROLE_SHORT.get(callers[-1]["role"], "") if callers else "",
                      "since": max((w["date"] for w in ws if w["date"]), default="")})
    cards.sort(key=lambda c: (c["state"] == "ok", c["i"]))
    open_st = request.args.get("open", type=int)
    measures = db.q("SELECT m.*, u.name AS uname FROM measures m LEFT JOIN users u ON u.id=m.user_id WHERE m.flat_id=? ORDER BY m.id DESC", (fid,))
    return render_template("flat.html", f=f, c=c, st=st, cards=cards, pics=pics, hist=hist, measures=measures,
                           open_st=open_st, tab=request.args.get("tab", "stages"), nav="home")


def photo_editable(w):
    owner = TPL[w["stage"]][1][w["idx"]][3]
    if owner == "okk":
        return can("okk") and w["status"] != "ok"
    return can("pr") and w["status"] in ("no", "wk")


app.jinja_env.globals["photo_editable"] = lambda w: g.user and photo_editable(w)


def work_or_404(wid):
    w = db.q("SELECT * FROM works WHERE id=?", (wid,), one=True)
    if not w:
        abort(404)
    flat_or_404(w["flat_id"])
    return w


@app.route("/w/<int:wid>/<action>", methods=["POST"])
@login_required
def work_action(wid, action):
    w = work_or_404(wid)
    name, _, need, owner = TPL[w["stage"]][1][w["idx"]]
    d = now()
    back = redirect(url_for("flat", fid=w["flat_id"], open=w["stage"]) + f"#s{w['stage']}")
    if owner == "okk":
        # дефектный акт: ОКК загружает фото и закрывает сам
        if action != "send" or not can("okk") or w["status"] == "ok":
            abort(403)
        n = db.q("SELECT COUNT(*) AS n FROM photos WHERE work_id=?", (wid,), one=True)["n"]
        if n < need:
            flash("Загрузите фото дефектного акта")
            return back
        db.x("UPDATE works SET status='ok', date=?, ret='', sent_by=? WHERE id=?", (d, g.user["id"], wid))
        log(w["flat_id"], w["stage"], "g", f"{name}: загружен ({n} фото), объект принят от заказчика")
        flash("Дефектный акт загружен: +5,56%")
        db.commit()
        return back
    if action == "start" and can("pr") and w["status"] == "no":
        db.x("UPDATE works SET status='wk', date=? WHERE id=?", (d, wid))
        log(w["flat_id"], w["stage"], "b", f"{name}: работы начаты")
    elif action == "send" and can("pr") and w["status"] in ("no", "wk"):
        n = db.q("SELECT COUNT(*) AS n FROM photos WHERE work_id=?", (wid,), one=True)["n"]
        if n < need:
            flash(f"Нужно минимум {need} фото, загружено {n}")
            return back
        db.x("UPDATE works SET status='rv', date=?, ret='', sent_by=? WHERE id=?", (d, g.user["id"], wid))
        log(w["flat_id"], w["stage"], "b", f"{name}: " + (f"фотоотчёт сдан ({n} фото)" if need else "сдано на проверку"))
        flash("Отправлено на проверку ОКК")
    elif action == "accept" and can("okk") and w["status"] == "rv":
        db.x("UPDATE works SET status='ok', date=?, ret='' WHERE id=?", (d, wid))
        log(w["flat_id"], w["stage"], "g", f"{name}: принято ОКК")
        rest = db.q("SELECT status FROM works WHERE flat_id=? AND stage=?", (w["flat_id"], w["stage"]))
        flash("Этап принят: +5,56%" if all(r["status"] == "ok" for r in rest) else "Работа принята")
    elif action == "return" and can("okk") and w["status"] == "rv":
        reason = request.form.get("reason", "").strip()[:200] or "требуется доработка"
        db.x("UPDATE works SET status='wk', ret=? WHERE id=?", (reason, wid))
        if need:
            db.x("DELETE FROM photos WHERE work_id=?", (wid,))
        log(w["flat_id"], w["stage"], "r", f"{name}: возвращено — {reason}")
        flash("Возвращено прорабу")
    else:
        abort(403)
    db.commit()
    return back


@app.route("/w/<int:wid>/photo", methods=["POST"])
@login_required
def work_photo(wid):
    w = work_or_404(wid)
    if not photo_editable(w):
        abort(403)
    from PIL import Image, ImageOps
    added = 0
    for f in request.files.getlist("photos")[:10]:
        if not f or not f.filename:
            continue
        try:
            im = ImageOps.exif_transpose(Image.open(f.stream)).convert("RGB")
            im.thumbnail((1600, 1600))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=80)
        except Exception:
            flash(f"Файл {f.filename} не похож на фото")
            continue
        db.x("INSERT INTO photos(work_id,data,mime,user_id,created) VALUES(?,?,?,?,?)",
             (wid, db.blob(buf.getvalue()), "image/jpeg", g.user["id"], now()))
        added += 1
    if added and w["status"] == "no":
        db.x("UPDATE works SET status='wk', date=? WHERE id=?", (now(), wid))
    db.commit()
    if added:
        flash(f"Фото добавлено: {added}")
    return redirect(url_for("flat", fid=w["flat_id"], open=w["stage"]) + f"#s{w['stage']}")


@app.route("/photo/<int:pid>")
@login_required
def photo(pid):
    p = db.q("SELECT * FROM photos WHERE id=?", (pid,), one=True)
    if not p:
        abort(404)
    work_or_404(p["work_id"])
    data = p["data"]
    return send_file(io.BytesIO(bytes(data)), mimetype=p["mime"], max_age=86400)


@app.route("/photo/<int:pid>/delete", methods=["POST"])
@login_required
def photo_delete(pid):
    p = db.q("SELECT * FROM photos WHERE id=?", (pid,), one=True)
    if not p:
        abort(404)
    w = work_or_404(p["work_id"])
    if not photo_editable(w):
        abort(403)
    db.x("DELETE FROM photos WHERE id=?", (pid,))
    db.commit()
    return redirect(url_for("flat", fid=w["flat_id"], open=w["stage"]) + f"#s{w['stage']}")


# ---------- этапы (сводка), задачи, уведомления ----------
@app.route("/stages")
@staff_required
def stages():
    ids = complex_ids()
    cs = db.q(f"SELECT * FROM complexes WHERE id IN {in_list(ids)} ORDER BY id")
    cid = request.args.get("c", type=int) or (cs[0]["id"] if cs else None)
    rows, flats_l = [], []
    if cid:
        check_complex(cid)
        flats_l, _ = complex_stats(cid)
        n = len(flats_l) or 1
        for i in range(N):
            o = sum(1 for f in flats_l if f["states"][i] == "ok")
            p = sum(1 for f in flats_l if f["sent"][i])
            rows.append({"i": i, "o": o, "p": p, "okk": o * 100 / n, "pr": p * 100 / n})
    measures = db.q("SELECT m.*, f.number, c.name AS cname FROM measures m JOIN flats f ON f.id=m.flat_id JOIN complexes c ON c.id=f.complex_id "
                    f"WHERE f.complex_id IN {in_list(ids)} AND m.status<>'done' ORDER BY m.id DESC")
    return render_template("stages.html", cs=cs, cid=cid, rows=rows, flats=flats_l, tasks=task_works(), measures=measures,
                           open_st=request.args.get("s", type=int), nav="stages")


@app.route("/tasks")
@login_required
def tasks():
    return redirect(url_for("stages") + "#tasks")


@app.route("/notifications")
@login_required
def notifications():
    ids = complex_ids()
    extra = f" AND f.id IN {in_list(client_flat_ids())}" if is_client() else ""
    evs = db.q("SELECT e.*, f.number, c.name AS cname, u.name AS uname FROM events e JOIN flats f ON f.id=e.flat_id "
               f"JOIN complexes c ON c.id=f.complex_id LEFT JOIN users u ON u.id=e.user_id WHERE f.complex_id IN {in_list(ids)}{extra} "
               "ORDER BY e.id DESC LIMIT 100")
    seen = g.user["notes_seen"] or ""
    db.x("UPDATE users SET notes_seen=? WHERE id=?", (now(), g.user["id"]))
    db.commit()
    return render_template("notifications.html", evs=evs, seen=seen, nav="home")


# ---------- профиль ----------
@app.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    if request.method == "POST":
        old, new = request.form.get("old", ""), request.form.get("new", "")
        if not check_password_hash(g.user["pw"], old):
            flash("Текущий пароль неверный")
        elif len(new) < 6:
            flash("Новый пароль — минимум 6 символов")
        else:
            db.x("UPDATE users SET pw=? WHERE id=?", (generate_password_hash(new), g.user["id"]))
            db.commit()
            flash("Пароль изменён")
        return redirect(url_for("profile"))
    n_c = len(complex_ids())
    return render_template("profile.html", n_c=n_c, nav="profile")


# ---------- администрирование ----------
@app.route("/admin/users", methods=["GET", "POST"])
@admin_required
def admin_users():
    if request.method == "POST":
        uid = request.form.get("uid", type=int)
        u = db.q("SELECT * FROM users WHERE id=?", (uid,), one=True)
        if not u or u["id"] == g.user["id"] and request.form.get("act") in ("block", "save", "approve"):
            flash("Свою учётную запись так менять нельзя")
            return redirect(url_for("admin_users"))
        act = request.form.get("act")
        if act in ("approve", "save"):
            role = request.form.get("role", u["role"])
            if role not in ROLES:
                abort(400)
            db.x("UPDATE users SET role=?, status='active' WHERE id=?", (role, uid))
            db.x("DELETE FROM access WHERE user_id=?", (uid,))
            db.x("DELETE FROM flat_access WHERE user_id=?", (uid,))
            if role == "client":
                for fid in request.form.getlist("fx", type=int):
                    db.x("INSERT INTO flat_access(user_id, flat_id) VALUES(?,?)", (uid, fid))
            else:
                for cid in request.form.getlist("cx", type=int):
                    db.x("INSERT INTO access(user_id, complex_id) VALUES(?,?)", (uid, cid))
            db.x("UPDATE users SET can_sched=? WHERE id=?", (1 if request.form.get("sched") else 0, uid))
            flash(f"{u['name']}: доступ сохранён")
        elif act == "block":
            db.x("UPDATE users SET status='blocked' WHERE id=?", (uid,))
            flash(f"{u['name']}: заблокирован")
        elif act == "unblock":
            db.x("UPDATE users SET status='active' WHERE id=?", (uid,))
        elif act == "reset":
            tmp = f"{secrets.randbelow(900000) + 100000}"
            db.x("UPDATE users SET pw=? WHERE id=?", (generate_password_hash(tmp), uid))
            flash(f"{u['name']}: временный пароль {tmp} — передайте сотруднику")
        db.commit()
        return redirect(url_for("admin_users"))
    users = db.q("SELECT * FROM users ORDER BY CASE status WHEN 'pending' THEN 0 WHEN 'active' THEN 1 ELSE 2 END, id DESC")
    acc = {}
    for a in db.q("SELECT * FROM access"):
        acc.setdefault(a["user_id"], set()).add(a["complex_id"])
    cs = db.q("SELECT * FROM complexes ORDER BY id")
    facc = {}
    for a in db.q("SELECT * FROM flat_access"):
        facc.setdefault(a["user_id"], set()).add(a["flat_id"])
    allflats = {}
    for f in db.q("SELECT * FROM flats ORDER BY id"):
        allflats.setdefault(f["complex_id"], []).append(f)
    return render_template("admin_users.html", users=users, acc=acc, facc=facc, allflats=allflats, cs=cs, nav="profile")


def parse_numbers(spec):
    out = []
    for part in re.split(r"[,\s;]+", spec or ""):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if 0 < b - a < 1000:
                out += [str(i) for i in range(a, b + 1)]
        else:
            out.append(part[:12])
    seen, res = set(), []
    for n in out:
        if n not in seen:
            seen.add(n)
            res.append(n)
    return res[:1000]


def add_flats(cid, numbers, kind):
    for n in numbers:
        fid = db.x("INSERT INTO flats(complex_id, number, kind) VALUES(?,?,?)", (cid, n, kind), returning=True)
        for si, (_, works, _) in enumerate(TPL):
            for wi in range(len(works)):
                db.x("INSERT INTO works(flat_id, stage, idx) VALUES(?,?,?)", (fid, si, wi))


@app.route("/admin/objects", methods=["GET", "POST"])
@admin_required
def admin_objects():
    if request.method == "POST":
        act = request.form.get("act")
        if act == "create":
            name = request.form.get("name", "").strip()[:120]
            nums = parse_numbers(request.form.get("numbers"))
            if not name or not nums:
                flash("Укажите название ЖК и номера квартир, например 1-42 или 187, 193, 199")
                return redirect(url_for("admin_objects"))
            cid = db.x("INSERT INTO complexes(name, sub, city, created) VALUES(?,?,?,?)",
                       (name, request.form.get("sub", "").strip()[:60], request.form.get("city", "Алматы").strip()[:60] or "Алматы", now()),
                       returning=True)
            add_flats(cid, nums, request.form.get("kind", "типовой"))
            flash(f"Создан ЖК «{name}»: {len(nums)} кв.")
        elif act == "add":
            cid = request.form.get("cid", type=int)
            nums = parse_numbers(request.form.get("numbers"))
            exist = {r["number"] for r in db.q("SELECT number FROM flats WHERE complex_id=?", (cid,))}
            nums = [n for n in nums if n not in exist]
            add_flats(cid, nums, request.form.get("kind", "типовой"))
            flash(f"Добавлено квартир: {len(nums)}")
        elif act == "delete":
            cid = request.form.get("cid", type=int)
            db.x("DELETE FROM complexes WHERE id=?", (cid,))
            flash("ЖК удалён")
        db.commit()
        return redirect(url_for("admin_objects"))
    cs = db.q("SELECT c.*, (SELECT COUNT(*) FROM flats f WHERE f.complex_id=c.id) AS n FROM complexes c ORDER BY c.id")
    return render_template("admin_objects.html", cs=cs, nav="profile")


# ---------- замеры ----------
@app.route("/f/<int:fid>/measure", methods=["POST"])
@staff_required
def measure_new(fid):
    f = flat_or_404(fid)
    kind = request.form.get("kind")
    if kind not in MEASURE_KINDS:
        abort(400)
    db.x("INSERT INTO measures(flat_id, kind, want_date, comment, user_id, created) VALUES(?,?,?,?,?,?)",
         (fid, kind, request.form.get("want_date", "")[:10], request.form.get("comment", "").strip()[:300], g.user["id"], now()))
    log(fid, None, "b", f"Вызван замер: {kind}")
    db.commit()
    flash(f"Замер вызван: {kind}")
    return redirect(url_for("flat", fid=f["id"]) + "#measures")


@app.route("/m/<int:mid>/<act>", methods=["POST"])
@staff_required
def measure_act(mid, act):
    m = db.q("SELECT * FROM measures WHERE id=?", (mid,), one=True)
    if not m:
        abort(404)
    flat_or_404(m["flat_id"])
    if g.user["role"] not in MEASURE_ROLES:
        abort(403)
    if act == "schedule":
        d = request.form.get("sched_date", "")[:10]
        db.x("UPDATE measures SET status='scheduled', sched_date=? WHERE id=?", (d, mid))
        log(m["flat_id"], None, "b", f"Замер «{m['kind']}» назначен на {ru_date(d)}")
    elif act == "done":
        db.x("UPDATE measures SET status='done' WHERE id=?", (mid,))
        log(m["flat_id"], None, "g", f"Замер «{m['kind']}» выполнен")
    elif act == "cancel":
        db.x("DELETE FROM measures WHERE id=?", (mid,))
    else:
        abort(400)
    db.commit()
    return redirect(request.form.get("back") or url_for("flat", fid=m["flat_id"]) + "#measures")


# ---------- финансы: общие расчёты ----------
def fact_rows(cid):
    """Факт расходов: оплаченные заявки + оплаченные строки графика выплат."""
    rows = db.q("SELECT flat_id, category, amount FROM requests WHERE complex_id=? AND status='paid'", (cid,))
    rows += db.q("SELECT flat_id, category, amount FROM schedule WHERE complex_id=? AND kind='pay' AND done=1", (cid,))
    return rows


def sum_by(rows, key):
    out = {}
    for r in rows:
        k = key(r)
        out[k] = out.get(k, 0) + (r["amount"] or 0)
    return out


def complex_limits(cid):
    lim = {(r["flat_id"] or 0): r for r in db.q("SELECT * FROM limits WHERE complex_id=?", (cid,))}
    return lim


def limit_check(cid, flat_id, category, extra=0):
    """Возвращает текст предупреждения, если лимит превышен."""
    if category not in ("rough", "finish"):
        return ""
    lim = complex_limits(cid)
    facts = fact_rows(cid)
    msgs = []
    for key, label in ((0, "блок"), (flat_id or -1, "квартиру")):
        L = lim.get(key)
        if not L or not L[category]:
            continue
        used = sum(r["amount"] for r in facts if r["category"] == category and (key == 0 or r["flat_id"] == key)) + extra
        if used > L[category]:
            msgs.append(f"превышен лимит «{CATS[category]}» на {label}: {money(used)} из {money(L[category])}")
    return "; ".join(msgs)


def gd_complexes():
    return db.q("SELECT * FROM complexes ORDER BY id")


# ---------- заявки на выплату / закуп ----------
@app.route("/requests", methods=["GET", "POST"])
@staff_required
def requests_page():
    ids = complex_ids()
    if request.method == "POST":
        cid = request.form.get("complex_id", type=int)
        if cid not in ids:
            abort(403)
        fid = request.form.get("flat_id", type=int) or None
        cat = request.form.get("category")
        kind = request.form.get("kind")
        try:
            amount = float(request.form.get("amount", "0").replace(" ", "").replace(",", "."))
        except ValueError:
            amount = 0
        if cat not in CATS or kind not in ("pay", "buy") or amount <= 0:
            flash("Заполните тип, категорию и сумму")
            return redirect(url_for("requests_page"))
        db.x("INSERT INTO requests(kind, complex_id, flat_id, category, amount, party, descr, user_id, created) VALUES(?,?,?,?,?,?,?,?,?)",
             (kind, cid, fid, cat, amount, request.form.get("party", "").strip()[:120], request.form.get("descr", "").strip()[:500], g.user["id"], now()))
        db.commit()
        warn = limit_check(cid, fid, cat, amount)
        flash("Заявка создана и отправлена генеральному директору" + (f". Внимание: {warn}" if warn else ""))
        return redirect(url_for("requests_page"))
    base = ("SELECT r.*, c.name AS cname, f.number, u.name AS uname FROM requests r JOIN complexes c ON c.id=r.complex_id "
            "LEFT JOIN flats f ON f.id=r.flat_id LEFT JOIN users u ON u.id=r.user_id ")
    if is_gd():
        reqs = db.q(base + "ORDER BY CASE r.status WHEN 'new' THEN 0 WHEN 'approved' THEN 1 ELSE 2 END, r.id DESC")
    else:
        reqs = db.q(base + "WHERE r.user_id=? ORDER BY r.id DESC", (g.user["id"],))
    cs = db.q(f"SELECT * FROM complexes WHERE id IN {in_list(ids)} ORDER BY id")
    flats_by = {}
    for f in db.q(f"SELECT * FROM flats WHERE complex_id IN {in_list(ids)} ORDER BY id"):
        flats_by.setdefault(f["complex_id"], []).append(f)
    return render_template("requests.html", reqs=reqs, cs=cs, flats_by=flats_by, nav="fin" if is_gd() else "profile",
                           pre_c=request.args.get("c", type=int), pre_f=request.args.get("f", type=int))


@app.route("/requests/<int:rid>/<act>", methods=["POST"])
@admin_required
def request_act(rid, act):
    st = {"approve": "approved", "reject": "rejected", "paid": "paid"}.get(act)
    if not st:
        abort(400)
    r = db.q("SELECT * FROM requests WHERE id=?", (rid,), one=True)
    if not r:
        abort(404)
    db.x("UPDATE requests SET status=?, decided=? WHERE id=?", (st, now(), rid))
    db.commit()
    if st == "paid":
        warn = limit_check(r["complex_id"], r["flat_id"], r["category"])
        flash("Отмечено как оплачено — сумма ушла в факт" + (f". Внимание: {warn}" if warn else ""))
    return redirect(url_for("requests_page"))


# ---------- финансы (только ген. директор) ----------
@app.route("/finance")
@login_required
def finance():
    if not is_gd():
        if sched_access():
            return redirect(url_for("schedule_page"))
        abort(403)
    cs = gd_complexes()
    t = request.args.get("t", "costs")
    cid = request.args.get("c", type=int) or (cs[0]["id"] if cs else None)
    data = {}
    if cid:
        c = db.q("SELECT * FROM complexes WHERE id=?", (cid,), one=True)
        flats_l = db.q("SELECT * FROM flats WHERE complex_id=? ORDER BY id", (cid,))
        facts = fact_rows(cid)
        plans = db.q("SELECT * FROM plans WHERE complex_id=?", (cid,))
        data = {
            "c": c, "flats": flats_l,
            "fact_cat": sum_by(facts, lambda r: r["category"]),
            "fact_flat": sum_by(facts, lambda r: r["flat_id"] or 0),
            "fact_flat_cat": sum_by(facts, lambda r: (r["flat_id"] or 0, r["category"])),
            "plan_cat": sum_by(plans, lambda r: r["category"]),
            "plan_flat_cat": sum_by(plans, lambda r: (r["flat_id"] or 0, r["category"])),
            "plans": plans, "limits": complex_limits(cid),
        }
        data["fact_total"] = sum(data["fact_cat"].values())
        data["plan_total"] = sum(data["plan_cat"].values())
    overview = []
    for c in cs:
        f = fact_rows(c["id"])
        p = db.q("SELECT COALESCE(SUM(amount),0) AS s FROM plans WHERE complex_id=?", (c["id"],), one=True)["s"]
        n = db.q("SELECT COUNT(*) AS n FROM flats WHERE complex_id=?", (c["id"],), one=True)["n"]
        overview.append({**c, "fact": sum(r["amount"] for r in f), "plan": p, "n": n})
    return render_template("finance.html", cs=cs, cid=cid, t=t, d=data, overview=overview, nav="fin")


@app.route("/finance/plan", methods=["POST"])
@admin_required
def finance_plan():
    cid = request.form.get("cid", type=int)
    if not db.q("SELECT id FROM complexes WHERE id=?", (cid,), one=True):
        abort(404)
    numbers = {r["number"]: r["id"] for r in db.q("SELECT id, number FROM flats WHERE complex_id=?", (cid,))}
    rows = []
    f = request.files.get("xlsx")
    if f and f.filename:
        from openpyxl import load_workbook
        try:
            ws = load_workbook(f.stream, data_only=True, read_only=True).active
            back_cats = {v.lower(): k for k, v in CATS.items()}
            for r in list(ws.iter_rows(values_only=True))[1:]:
                if not r or r[2] in (None, ""):
                    continue
                num = str(r[0]).strip() if r[0] not in (None, "") else ""
                cat = back_cats.get(str(r[1] or "").strip().lower())
                if not cat:
                    continue
                rows.append((numbers.get(num) if num else None, cat, float(str(r[2]).replace(" ", "").replace(",", "."))))
        except Exception as e:  # noqa: BLE001
            flash(f"Не удалось прочитать Excel: {e}")
            return redirect(url_for("finance", t="planfact", c=cid))
        db.x("DELETE FROM plans WHERE complex_id=?", (cid,))
    else:
        cat = request.form.get("category")
        fid = request.form.get("flat_id", type=int) or None
        try:
            amount = float(request.form.get("amount", "0").replace(" ", "").replace(",", "."))
        except ValueError:
            amount = 0
        if cat in CATS and amount >= 0:
            if fid:
                db.x("DELETE FROM plans WHERE complex_id=? AND flat_id=? AND category=?", (cid, fid, cat))
            else:
                db.x("DELETE FROM plans WHERE complex_id=? AND flat_id IS NULL AND category=?", (cid, cat))
            rows.append((fid, cat, amount))
    for fid, cat, amount in rows:
        db.x("INSERT INTO plans(complex_id, flat_id, category, amount) VALUES(?,?,?,?)", (cid, fid, cat, amount))
    db.commit()
    flash(f"План обновлён: строк {len(rows)}")
    return redirect(url_for("finance", t="planfact", c=cid))


@app.route("/finance/plan-template.xlsx")
@admin_required
def plan_template():
    from openpyxl import Workbook
    cid = request.args.get("c", type=int)
    wb = Workbook()
    ws = wb.active
    ws.title = "План"
    ws.append(["Квартира (пусто = весь блок)", "Категория", "Сумма, ₸"])
    for f in db.q("SELECT number FROM flats WHERE complex_id=? ORDER BY id", (cid,)):
        for k in ("rough", "finish", "works"):
            ws.append([f["number"], CATS[k], None])
    for col, w in zip("ABC", (28, 22, 16)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="plan.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/finance/limits", methods=["POST"])
@admin_required
def finance_limits():
    cid = request.form.get("cid", type=int)
    fid = request.form.get("flat_id", type=int) or None

    def num(k):
        try:
            return float(request.form.get(k, "0").replace(" ", "").replace(",", ".") or 0)
        except ValueError:
            return 0
    if fid:
        db.x("DELETE FROM limits WHERE complex_id=? AND flat_id=?", (cid, fid))
    else:
        db.x("DELETE FROM limits WHERE complex_id=? AND flat_id IS NULL", (cid,))
    db.x("INSERT INTO limits(complex_id, flat_id, rough, finish) VALUES(?,?,?,?)", (cid, fid, num("rough"), num("finish")))
    db.commit()
    flash("Лимит сохранён")
    return redirect(url_for("finance", t="limits", c=cid))


# ---------- график работ и выплат ----------
@app.route("/finance/schedule", methods=["GET", "POST"])
@login_required
def schedule_page():
    if not sched_access():
        abort(403)
    cs = gd_complexes() if is_gd() else db.q(f"SELECT * FROM complexes WHERE id IN {in_list(complex_ids())} ORDER BY id")
    allowed = {c["id"] for c in cs}
    if request.method == "POST":
        act = request.form.get("act")
        if act == "create":
            cid = request.form.get("complex_id", type=int)
            if cid not in allowed:
                abort(403)
            try:
                amount = float(request.form.get("amount", "0").replace(" ", "").replace(",", ".") or 0)
            except ValueError:
                amount = 0
            kind = request.form.get("kind") if request.form.get("kind") in ("work", "pay") else "pay"
            cat = request.form.get("category") if request.form.get("category") in CATS else "works"
            db.x("INSERT INTO schedule(complex_id, flat_id, kind, title, party, day, amount, category, user_id, created) VALUES(?,?,?,?,?,?,?,?,?,?)",
                 (cid, request.form.get("flat_id", type=int) or None, kind, request.form.get("title", "").strip()[:200] or "Без названия",
                  request.form.get("party", "").strip()[:120], request.form.get("day", "")[:10] or now()[:10], amount, cat, g.user["id"], now()))
            flash("Строка добавлена в график")
        else:
            sid = request.form.get("sid", type=int)
            row = db.q("SELECT * FROM schedule WHERE id=?", (sid,), one=True)
            if not row or row["complex_id"] not in allowed:
                abort(404)
            if act == "toggle":
                db.x("UPDATE schedule SET done=? WHERE id=?", (0 if row["done"] else 1, sid))
            elif act == "delete":
                db.x("DELETE FROM schedule WHERE id=?", (sid,))
        db.commit()
        return redirect(url_for("schedule_page", c=request.form.get("c") or None))
    cid = request.args.get("c", type=int) or (cs[0]["id"] if cs else None)
    rows = db.q("SELECT s.*, f.number FROM schedule s LEFT JOIN flats f ON f.id=s.flat_id WHERE s.complex_id=? ORDER BY s.day, s.id", (cid,)) if cid else []
    today = now()[:10]
    for r in rows:
        r["carry"] = (not r["done"]) and r["day"] < today
    tot = {
        "pay_all": sum(r["amount"] for r in rows if r["kind"] == "pay"),
        "pay_left": sum(r["amount"] for r in rows if r["kind"] == "pay" and not r["done"]),
        "work_all": sum(r["amount"] for r in rows if r["kind"] == "work"),
        "work_left": sum(r["amount"] for r in rows if r["kind"] == "work" and not r["done"]),
        "carry": sum(r["amount"] for r in rows if r["carry"]),
    }
    flats_l = db.q("SELECT * FROM flats WHERE complex_id=? ORDER BY id", (cid,)) if cid else []
    return render_template("schedule.html", cs=cs, cid=cid, rows=rows, tot=tot, flats=flats_l, today=today, nav="fin")


# ---------- юридический отдел ----------
def legal_required(f):
    @wraps(f)
    @login_required
    def w(*a, **k):
        if g.user["role"] not in LEGAL_ROLES:
            abort(403)
        return f(*a, **k)
    return w


def contract_or_404(kid):
    k = db.q("SELECT * FROM contracts WHERE id=?", (kid,), one=True)
    if not k:
        abort(404)
    if not is_gd() and not db.q("SELECT 1 AS x FROM contract_access WHERE contract_id=? AND user_id=?", (kid, g.user["id"]), one=True):
        abort(404)
    return k


@app.route("/legal", methods=["GET", "POST"])
@legal_required
def legal():
    kind = request.args.get("k", "client")
    if kind not in CONTRACT_KINDS:
        kind = "client"
    if request.method == "POST":
        if g.user["role"] not in ("admin", "lawyer"):
            abort(403)
        try:
            amount = float(request.form.get("amount", "0").replace(" ", "").replace(",", ".") or 0)
        except ValueError:
            amount = 0
        kid = db.x("INSERT INTO contracts(kind, title, party, number, cdate, amount, complex_id, note, user_id, created) VALUES(?,?,?,?,?,?,?,?,?,?)",
                   (kind, request.form.get("title", "").strip()[:200] or "Договор", request.form.get("party", "").strip()[:200],
                    request.form.get("number", "").strip()[:60], request.form.get("cdate", "")[:10], amount,
                    request.form.get("complex_id", type=int) or None, request.form.get("note", "").strip()[:500], g.user["id"], now()),
                   returning=True)
        if not is_gd():
            db.x("INSERT INTO contract_access(contract_id, user_id) VALUES(?,?)", (kid, g.user["id"]))
        db.commit()
        flash("Договор создан. Доступ сотрудникам открывает генеральный директор.")
        return redirect(url_for("contract", kid=kid))
    base = "SELECT k.*, c.name AS cname, (SELECT COUNT(*) FROM contract_files x WHERE x.contract_id=k.id) AS nfiles FROM contracts k LEFT JOIN complexes c ON c.id=k.complex_id "
    if is_gd():
        ks = db.q(base + "WHERE k.kind=? ORDER BY k.id DESC", (kind,))
    else:
        ks = db.q(base + "JOIN contract_access a ON a.contract_id=k.id AND a.user_id=? WHERE k.kind=? ORDER BY k.id DESC", (g.user["id"], kind))
    return render_template("legal.html", ks=ks, kind=kind, cs=gd_complexes(), nav="legal")


@app.route("/legal/<int:kid>", methods=["GET", "POST"])
@legal_required
def contract(kid):
    k = contract_or_404(kid)
    if request.method == "POST":
        act = request.form.get("act")
        if act == "upload":
            if g.user["role"] not in ("admin", "lawyer"):
                abort(403)
            n = 0
            for f in request.files.getlist("files")[:10]:
                if f and f.filename:
                    data = f.read()
                    if len(data) > 15 * 1024 * 1024:
                        flash(f"{f.filename}: больше 15 МБ")
                        continue
                    db.x("INSERT INTO contract_files(contract_id, name, mime, data, created) VALUES(?,?,?,?,?)",
                         (kid, f.filename[:150], f.mimetype or "application/octet-stream", db.blob(data), now()))
                    n += 1
            flash(f"Файлов загружено: {n}")
        elif act == "access":
            if not is_gd():
                abort(403)
            db.x("DELETE FROM contract_access WHERE contract_id=?", (kid,))
            for uid in request.form.getlist("uid", type=int):
                db.x("INSERT INTO contract_access(contract_id, user_id) VALUES(?,?)", (kid, uid))
            flash("Доступ к договору обновлён")
        elif act == "delete":
            if not is_gd():
                abort(403)
            db.x("DELETE FROM contracts WHERE id=?", (kid,))
            db.commit()
            flash("Договор удалён")
            return redirect(url_for("legal", k=k["kind"]))
        db.commit()
        return redirect(url_for("contract", kid=kid))
    files = db.q("SELECT id, name, created FROM contract_files WHERE contract_id=? ORDER BY id", (kid,))
    people = db.q("SELECT id, name, role FROM users WHERE status='active' AND role IN ('devdir','lawyer','accountant') ORDER BY name")
    has = {r["user_id"] for r in db.q("SELECT user_id FROM contract_access WHERE contract_id=?", (kid,))}
    c = db.q("SELECT name FROM complexes WHERE id=?", (k["complex_id"],), one=True) if k["complex_id"] else None
    return render_template("contract.html", k=k, files=files, people=people, has=has, c=c, nav="legal")


@app.route("/legal/file/<int:fid>")
@legal_required
def contract_file(fid):
    f = db.q("SELECT * FROM contract_files WHERE id=?", (fid,), one=True)
    if not f:
        abort(404)
    contract_or_404(f["contract_id"])
    return send_file(io.BytesIO(bytes(f["data"])), mimetype=f["mime"], download_name=f["name"])


@app.route("/sw.js")
def sw():
    return app.send_static_file("sw.js"), 200, {"Content-Type": "application/javascript", "Service-Worker-Allowed": "/"}


@app.errorhandler(403)
def e403(_e):
    return render_template("error.html", msg="Это действие недоступно для вашей роли"), 403


@app.errorhandler(404)
def e404(_e):
    return render_template("error.html", msg="Страница не найдена или нет доступа"), 404


if __name__ == "__main__":
    app.run(debug=True)
